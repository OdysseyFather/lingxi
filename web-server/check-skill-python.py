"""Offline smoke check for the Python used by Web OnCall skills."""
import io
import json
import sys

import requests
from openpyxl import Workbook, load_workbook


def main():
    request = requests.Request("GET", "https://example.invalid", params={"page": 1}).prepare()
    if request.url != "https://example.invalid/?page=1":
        raise RuntimeError("requests preparation failed")
    output = io.BytesIO()
    workbook = Workbook()
    workbook.active.append(["event", "count"])
    workbook.active.append(["payment_click", 1])
    workbook.save(output)
    workbook.close()
    output.seek(0)
    restored = load_workbook(output, read_only=True)
    try:
        if list(restored.active.values) != [("event", "count"), ("payment_click", 1)]:
            raise RuntimeError("xlsx round-trip failed")
    finally:
        restored.close()
    print(json.dumps({"python": sys.executable, "requests": requests.__version__, "xlsx": "ok"}))


if __name__ == "__main__":
    main()
